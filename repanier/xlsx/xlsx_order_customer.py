"""
Génération des exports Excel liés aux clients, résumés, étiquettes et listes de préparation.
"""

from django.conf import settings
from django.utils import translation
from django.utils.translation import gettext_lazy as _
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles.numbers import FORMAT_DATE_DMYSLASH, FORMAT_TEXT
from openpyxl.utils import get_column_letter

import repanier.apps
from repanier.const import (
    DECIMAL_ZERO,
    EMPTY_STRING,
    OrderUnit,
    Placement,
    RoundUpTo, REPANIER_MONEY_ZERO,
)
from repanier.models.customer import Customer
from repanier.models.deliveryboard import DeliveryBoard
from repanier.models.invoice import CustomerInvoice, ProducerInvoice
from repanier.models.permanence import Permanence
from repanier.models.permanenceboard import PermanenceBoard
from repanier.models.producer import Producer
from repanier.models.purchase import Purchase
from repanier.models.staff import Staff
from repanier.tools import get_base_unit, next_row
from repanier.xlsx.export_tools import (
    new_portrait_a4_sheet,
    new_landscape_a4_sheet,
    format_worksheet_title,
)
from repanier.xlsx.xlsx_styles import (
    FILL_YELLOW,
    BORDER_BOTTOM_THIN,
    BORDER_BOTTOM_MEDIUM_DASHED,
    BORDER_LABEL_TOP,
    BORDER_LABEL_BOTTOM,
    FONT_BOLD,
    FONT_ITALIC,
    FONT_RED,
    FONT_GRAY,
    FONT_LABEL_BIG,
    ALIGN_CENTER,
    ALIGN_RIGHT,
    ALIGN_LEFT_WRAP,
    ALIGN_LEFT_NOWRAP, FONT_BOLD_ITALIC,
)


def next_purchase(purchases):
    """Retourne le prochain achat de l'itérateur."""
    return next_row(purchases)


def export_abstract(permanence, deliveries_id=(), group=False, wb=None):
    """
    Exporte la feuille récapitulative de la permanence (clients, permanenciers, staff, producteurs).
    """
    if permanence is not None:
        row_num = 1
        Customer.objects.all().update(preparation_order=0)
        preparation_order = 1

        if len(deliveries_id) > 0:
            header = [
                (_("Delivery point"), 25),
                (_("Family"), 30),
                (_("Phone1"), 15),
                (_("Phone2"), 15),
                (_("Total with vat"), 15),
                (_("Email"), 35),
            ]
            wb, ws = new_portrait_a4_sheet(
                wb, permanence, EMPTY_STRING, header=header, add_print_title=False
            )
            for delivery_ref, delivery in enumerate(
                DeliveryBoard.objects.filter(id__in=deliveries_id).order_by("id")
            ):
                customer_set = Customer.objects.filter(
                    customerinvoice__permanence_id=permanence.id,
                    represent_this_buyinggroup=False,
                    is_group=False,
                    customerinvoice__delivery_id=delivery.id,
                )
                for customer in customer_set:
                    invoice = CustomerInvoice.objects.filter(
                        permanence=permanence, customer=customer
                    ).first()
                    if invoice is not None and invoice.has_purchase:
                        Customer.objects.filter(id=customer.id).update(
                            preparation_order=preparation_order
                        )
                        preparation_order += 1
                        if (
                            settings.REPANIER_SETTINGS_CUSTOMER_MUST_CONFIRM_ORDER
                            and not invoice.is_order_confirm_send
                        ):
                            confirmed = "\n{}".format(_("⚠ This order isn't confirmed"))
                        else:
                            confirmed = EMPTY_STRING
                        row = [
                            "{} - {}".format(
                                delivery_ref, delivery.get_delivery_display()
                            ),
                            "  {} - {}{}".format(
                                customer.preparation_order,
                                customer.long_basket_name,
                                confirmed,
                            ),
                            customer.get_phone1(),
                            customer.get_phone2(),
                            invoice.get_total_price_with_tax().amount,
                            customer.get_email_postfixed(),
                        ]
                        for col_num in range(len(row)):
                            c = ws.cell(row=row_num + 1, column=col_num + 1)
                            c.value = row[col_num]
                            if col_num == 4:
                                c.number_format = (
                                    repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                                )
                            else:
                                c.number_format = FORMAT_TEXT
                                if col_num == 1:
                                    c.alignment = ALIGN_LEFT_WRAP
                                else:
                                    c.alignment = ALIGN_LEFT_NOWRAP
                            if row_num % 2 == 0:
                                c.border = BORDER_BOTTOM_THIN
                        row_num += 1
        else:
            header = [
                (_("Basket"), 20),
                (_("Family"), 35),
                (_("Phone1"), 15),
                (_("Phone2"), 15),
                (_("Total with vat"), 15),
                (_("Email"), 35),
            ]
            wb, ws = new_portrait_a4_sheet(
                wb, permanence, EMPTY_STRING, header=header, add_print_title=False
            )
            customer_set = Customer.objects.filter(
                customerinvoice__permanence_id=permanence.id,
                represent_this_buyinggroup=False,
            )
            for customer in customer_set:
                invoice = CustomerInvoice.objects.filter(
                    permanence=permanence, customer=customer
                ).first()
                if invoice is not None and invoice.has_purchase:
                    Customer.objects.filter(id=customer.id).update(
                        preparation_order=preparation_order
                    )
                    preparation_order += 1
                    if (
                        settings.REPANIER_SETTINGS_CUSTOMER_MUST_CONFIRM_ORDER
                        and not invoice.is_order_confirm_send
                    ):
                        confirmed = "\n{}".format(_("⚠ This order isn't confirmed"))
                    else:
                        confirmed = EMPTY_STRING
                    row = [
                        "{} - {}{}".format(
                            customer.preparation_order,
                            customer.short_basket_name,
                            confirmed,
                        ),
                        customer.long_basket_name,
                        customer.get_phone1(),
                        customer.get_phone2(),
                        invoice.get_total_price_with_tax().amount,
                        customer.get_email_postfixed(),
                    ]
                    for col_num in range(len(row)):
                        c = ws.cell(row=row_num + 1, column=col_num + 1)
                        c.value = row[col_num]
                        if col_num == 4:
                            c.number_format = (
                                repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                            )
                        else:
                            c.number_format = FORMAT_TEXT
                            if col_num == 0:
                                c.alignment = ALIGN_LEFT_WRAP
                            else:
                                c.alignment = ALIGN_LEFT_NOWRAP
                        if row_num % 2 == 0:
                            c.border = BORDER_BOTTOM_THIN
                    row_num += 1

        # Permanence board info
        permanence_date_save = None
        at_least_one = False
        next_permanence_set = Permanence.objects.filter(
            permanence_date__gte=permanence.permanence_date
        ).order_by("permanence_date")[:5]
        for next_permanence in next_permanence_set:
            for permanenceboard in PermanenceBoard.objects.filter(
                permanence_id=next_permanence.id
            ).order_by("permanence_role__tree_id", "permanence_role__lft"):
                customer = permanenceboard.customer
                if customer is not None:
                    if not at_least_one:
                        for col_num in range(6):
                            c = ws.cell(row=row_num + 1, column=col_num + 1)
                            c.border = BORDER_BOTTOM_THIN

                        c = ws.cell(row=row_num + 1, column=0 + 1)
                        c.value = "-------"
                        c = ws.cell(row=row_num + 1, column=1 + 1)
                        c.value = "{}".format(_("Permanence Board Member List"))
                        c.alignment = ALIGN_LEFT_NOWRAP
                        c.font = FONT_BOLD
                        c = ws.cell(row=row_num + 1, column=2 + 1)
                        c.value = "-------"
                        row_num += 1

                    at_least_one = True
                    row = [
                        (
                            next_permanence.permanence_date,
                            FORMAT_DATE_DMYSLASH,
                        ),
                        (customer.long_basket_name, FORMAT_TEXT),
                        (customer.get_phone1(), FORMAT_TEXT),
                        (customer.get_phone2(), FORMAT_TEXT),
                        (
                            permanenceboard.permanence_role.short_name_v2,
                            FORMAT_TEXT,
                        ),
                    ]
                    for col_num in range(len(row)):
                        c = ws.cell(row=row_num + 1, column=col_num + 1)
                        c.value = row[col_num][0]
                        c.number_format = row[col_num][1]
                        c.alignment = ALIGN_LEFT_NOWRAP
                        if permanence_date_save != next_permanence.permanence_date:
                            c.font = FONT_BOLD
                            permanence_date_save = next_permanence.permanence_date
                            for _col_num in range(6):
                                c = ws.cell(row=row_num + 1, column=_col_num + 1)
                                c.border = BORDER_BOTTOM_THIN
                    row_num += 1

        at_least_one = False
        for staff in Staff.objects.filter(is_active=True):
            customer_responsible = staff.customer_responsible
            if customer_responsible is not None:
                if not at_least_one:
                    for col_num in range(6):
                        c = ws.cell(row=row_num + 1, column=col_num + 1)
                        c.border = BORDER_BOTTOM_THIN

                    c = ws.cell(row=row_num + 1, column=0 + 1)
                    c.value = "-------"
                    c = ws.cell(row=row_num + 1, column=1 + 1)
                    c.value = "{}".format(_("Staff Member List"))
                    c.alignment = ALIGN_LEFT_NOWRAP
                    c.font = FONT_BOLD
                    c = ws.cell(row=row_num + 1, column=2 + 1)
                    c.value = "-------"
                    row_num += 1

                    at_least_one = True

                staff_function = staff.long_name_v2
                row = [
                    staff_function,
                    customer_responsible.long_basket_name,
                    customer_responsible.get_phone1(),
                ]
                for col_num in range(len(row)):
                    c = ws.cell(row=row_num + 1, column=col_num + 1)
                    c.value = row[col_num]
                    c.number_format = FORMAT_TEXT
                    c.alignment = ALIGN_LEFT_WRAP
                row_num += 1

        if not group:
            for col_num in range(6):
                c = ws.cell(row=row_num + 1, column=col_num + 1)
                c.border = BORDER_BOTTOM_THIN

            c = ws.cell(row=row_num + 1, column=0 + 1)
            c.value = "-------"
            c = ws.cell(row=row_num + 1, column=1 + 1)
            c.value = "{}".format(_("Producers"))
            c.alignment = ALIGN_LEFT_NOWRAP
            c.font = FONT_BOLD
            c = ws.cell(row=row_num + 1, column=2 + 1)
            c.value = "-------"
            row_num += 1

            for producer in Producer.objects.filter(permanence=permanence).order_by(
                "short_profile_name"
            ):
                invoice = ProducerInvoice.objects.filter(
                    permanence=permanence, producer=producer
                ).first()
                minimum_order_amount_reached = EMPTY_STRING
                if invoice is None:
                    total_price_with_tax = REPANIER_MONEY_ZERO
                else:
                    total_price_with_tax = invoice.get_total_price_with_tax()
                    if total_price_with_tax < producer.minimum_order_value:
                        minimum_order_amount_reached = _(
                            "Minimum order amount not reached"
                        )
                row = [
                    producer.short_profile_name,
                    producer.long_profile_name,
                    producer.get_phone1(),
                    producer.get_phone2(),
                    total_price_with_tax.amount,
                    minimum_order_amount_reached,
                ]
                for col_num in range(len(row)):
                    c = ws.cell(row=row_num + 1, column=col_num + 1)
                    c.value = row[col_num]
                    if col_num == 4:
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    else:
                        c.number_format = FORMAT_TEXT
                        c.alignment = ALIGN_LEFT_NOWRAP
                row_num += 1

        return wb

    return


def customer_label(customer_identifier, placements, row_num, ws):
    """Génère une étiquette individuelle pour un client."""
    label_row = row_num + 1
    c = ws.cell(row=label_row, column=1, value=str(customer_identifier))
    c.font = FONT_LABEL_BIG
    c.alignment = ALIGN_CENTER
    c.border = BORDER_LABEL_TOP
    ws.row_dimensions[label_row].height = 60

    row_num += 1
    c = ws.cell(
        row=row_num + 1,
        column=1,
        value=", ".join(map(str, placements)),
    )
    c.alignment = ALIGN_CENTER
    c.border = BORDER_LABEL_BOTTOM

    row_num += 1
    spacer_row = row_num + 1
    ws.row_dimensions[spacer_row].height = 5
    return spacer_row


def export_customer_label(permanence, deliveries_id=(), wb=None):
    """Exporte la feuille des étiquettes clients."""
    wb, ws = new_portrait_a4_sheet(wb, _("Label"), permanence, add_print_title=False)
    row_num = 0
    customer_set = Customer.objects.filter(
        customerinvoice__permanence_id=permanence.id,
        represent_this_buyinggroup=False,
        preparation_order__gt=0,
    ).order_by("preparation_order")

    if len(deliveries_id) > 0:
        customer_set = customer_set.filter(
            customerinvoice__delivery_id__in=deliveries_id
        )

    freezer = []
    fridge = []
    out_of_basket = []

    for customer in customer_set:
        customer_identifier = "{} - {}".format(
            customer.preparation_order, customer.short_basket_name
        )
        placements = [
            Placement(placement)
            for placement in list(
                Purchase.objects.filter(
                    permanence_id=permanence.id, customer_id=customer.id
                )
                .order_by("offer_item__placement")
                .values_list("offer_item__placement", flat=True)
                .distinct("offer_item__placement")
            )
        ]

        if Placement.FREEZER in placements:
            freezer += [customer_identifier]
        if Placement.FRIDGE in placements:
            fridge += [customer_identifier]
        if Placement.OUT_OF_BASKET in placements:
            out_of_basket += [customer_identifier]

        row_num = customer_label(
            customer_identifier,
            [placement.label for placement in placements],
            row_num,
            ws,
        )

    placement_label = "[❄ {} ]".format(Placement.FRIDGE.label)
    for customer_identifier in fridge:
        row_num = customer_label(
            customer_identifier,
            [
                placement_label,
            ],
            row_num,
            ws,
        )

    placement_label = "[❄❄❄ {} ]".format(Placement.FREEZER.label)
    for customer_identifier in freezer:
        row_num = customer_label(
            customer_identifier,
            [
                placement_label,
            ],
            row_num,
            ws,
        )

    placement_label = "👜 [ {} ]".format(Placement.OUT_OF_BASKET.label)
    for customer_identifier in out_of_basket:
        row_num = customer_label(
            customer_identifier,
            [
                placement_label,
            ],
            row_num,
            ws,
        )

    if row_num > 0:
        ws.column_dimensions[get_column_letter(1)].width = 120

    return wb


def export_preparation(permanence, deliveries_id=(), wb=None):
    """Exporte la liste de préparation des commandes."""
    header = [
        (_("Purchase"), 10),
        (_("OfferItem"), 5),
        (_("Placement"), 7),
        (_("Producer"), 15),
        (EMPTY_STRING, 7),
        (_("Product"), 60),
        (_("Customer"), 15),
        (_("Quantity ordered"), 10),
        (_("Unit"), 10),
        (_("Prepared"), 22),
        (_("To distribute"), 10),
    ]
    if len(deliveries_id) == 0:
        return export_preparation_for_a_delivery(
            0, None, header, permanence, wb, FILL_YELLOW
        )
    else:
        for delivery_ref, delivery_id in enumerate(deliveries_id):
            wb = export_preparation_for_a_delivery(
                delivery_ref, delivery_id, header, permanence, wb, FILL_YELLOW
            )
        return wb


def export_preparation_for_a_delivery(
    delivery_cpt, delivery_id, header, permanence, wb, yellow_fill
):
    """Génère la préparation pour une livraison donnée."""
    from repanier.models.offeritem import OfferItemReadOnly

    producer_set = Producer.objects.filter(
        producerinvoice__permanence_id=permanence.id
    ).only("invoice_by_basket", "short_profile_name")
    if delivery_id is not None:
        producer_set = producer_set.filter(
            purchase__customer_invoice__delivery_id=delivery_id
        )
    producers = producer_set.distinct().iterator()
    producer = next_row(producers)
    if producer is not None:
        wb, ws = new_landscape_a4_sheet(
            wb,
            _("Preparation")
            if delivery_id is None
            else "{}-{}".format(delivery_cpt, _("Preparation")),
            permanence,
            header,
        )
        row_num = 1
        if delivery_id is not None:
            c = ws.cell(row=row_num + 1, column=5 + 1)
            c.font = FONT_BOLD
            c.value = (
                DeliveryBoard.objects.filter(id=delivery_id)
                .first()
                .get_delivery_display()
            )
            row_num += 1
        producer_counter = 0
        hide_column_placement = True
        placement_save = None
        while producer is not None:
            producer_counter += 1
            producer_save = producer
            at_least_one_product = False

            if producer.invoice_by_basket and not producer.represent_this_buyinggroup:
                purchase_set = (
                    Purchase.objects.filter(
                        permanence_id=permanence.id,
                        producer_id=producer.id,
                    )
                    .order_by(
                        "customer__short_basket_name",
                        "offer_item__preparation_sort_order_v2",
                    )
                    .select_related(
                        "customer", "offer_item", "offer_item__department_for_customer"
                    )
                )
                if delivery_id is not None:
                    purchase_set = purchase_set.filter(
                        customer_invoice__delivery_id=delivery_id
                    )
                purchases = purchase_set.iterator()
                purchase = next_purchase(purchases)
                while purchase is not None:
                    at_least_one_product = True
                    customer_save = purchase.customer
                    c = ws.cell(row=row_num + 1, column=6 + 1)
                    c.font = FONT_BOLD
                    row_num += 1
                    count_purchase = 0
                    purchases_price = DECIMAL_ZERO
                    purchases_price_formula = []
                    while (
                        purchase is not None
                        and customer_save.id == purchase.customer_id
                    ):
                        department_for_customer_save = (
                            purchase.offer_item.department_for_customer
                        )
                        department_for_customer_save__short_name = (
                            department_for_customer_save.short_name_v2
                            if department_for_customer_save is not None
                            else EMPTY_STRING
                        )
                        while (
                            purchase is not None
                            and customer_save.id == purchase.customer_id
                            and department_for_customer_save
                            == purchase.offer_item.department_for_customer
                        ):
                            qty = purchase.get_producer_quantity()
                            if qty != DECIMAL_ZERO:
                                base_unit = get_base_unit(
                                    purchase.offer_item.order_unit, qty
                                )
                                c = ws.cell(row=row_num + 1, column=0 + 1)
                                c.value = purchase.id
                                c = ws.cell(row=row_num + 1, column=1 + 1)
                                c.value = purchase.offer_item_id
                                c.number_format = FORMAT_DATE_DMYSLASH
                                c = ws.cell(row=row_num + 1, column=2 + 1)
                                c.value = "{}".format(
                                    purchase.offer_item.get_placement_display()
                                )
                                if placement_save is None:
                                    placement_save = c.value
                                elif hide_column_placement:
                                    if placement_save != c.value:
                                        hide_column_placement = False
                                c = ws.cell(row=row_num + 1, column=3 + 1)
                                c.value = "{}".format(producer_save.short_profile_name)
                                c = ws.cell(row=row_num + 1, column=5 + 1)
                                if (
                                    department_for_customer_save__short_name
                                    is not None
                                ):
                                    c.value = "{} - {}".format(
                                        purchase.get_long_name_with_customer_price(),
                                        department_for_customer_save__short_name,
                                    )
                                else:
                                    c.value = "{}".format(
                                        purchase.get_long_name_with_customer_price()
                                    )
                                c.number_format = FORMAT_TEXT
                                c.alignment = ALIGN_LEFT_WRAP
                                c = ws.cell(row=row_num + 1, column=6 + 1)
                                c.value = "{} - {}".format(
                                    customer_save.preparation_order,
                                    customer_save.short_basket_name,
                                )
                                c.number_format = FORMAT_TEXT
                                c = ws.cell(row=row_num + 1, column=7 + 1)
                                c.value = qty
                                c.number_format = "#,##0.????"
                                c.font = FONT_RED
                                ws.conditional_formatting.add(
                                    f"{get_column_letter(8)}{row_num + 1}",
                                    CellIsRule(
                                        operator="notEqual",
                                        formula=[str(qty)],
                                        fill=yellow_fill,
                                    ),
                                )
                                c = ws.cell(row=row_num + 1, column=8 + 1)
                                c.value = "{}".format(base_unit)
                                c.number_format = FORMAT_TEXT
                                customer_unit_price = (
                                    purchase.get_customer_unit_price()
                                )
                                if purchase.offer_item.order_unit == OrderUnit.PC_KG:
                                    price_qty = (
                                        qty * purchase.offer_item.order_average_weight
                                    )
                                    purchases_price_formula.append(
                                        "ROUND(H{}*{}*{},2)".format(
                                            row_num + 1,
                                            purchase.offer_item.order_average_weight,
                                            customer_unit_price
                                            + purchase.offer_item.unit_deposit.amount,
                                        )
                                    )
                                else:
                                    price_qty = qty
                                    purchases_price_formula.append(
                                        "ROUND(H{}*{},2)".format(
                                            row_num + 1,
                                            customer_unit_price
                                            + purchase.offer_item.unit_deposit.amount,
                                        )
                                    )
                                if purchase.offer_item.order_unit == OrderUnit.PC_KG:
                                    c = ws.cell(row=row_num + 1, column=9 + 1)
                                    if purchase.offer_item.wrapped:
                                        c.value = "{} :".format(
                                            repanier.apps.REPANIER_SETTINGS_CURRENCY_DISPLAY
                                        )
                                    else:
                                        c.value = "{}".format(_("kg :"))
                                    c.number_format = FORMAT_TEXT
                                else:
                                    if purchase.offer_item.wrapped:
                                        c = ws.cell(row=row_num + 1, column=9 + 1)
                                        c.value = "{} :".format(
                                            repanier.apps.REPANIER_SETTINGS_CURRENCY_DISPLAY
                                        )
                                        c.number_format = FORMAT_TEXT
                                purchases_price += (
                                    price_qty
                                    * (
                                        customer_unit_price
                                        + purchase.offer_item.unit_deposit.amount
                                    )
                                ).quantize(RoundUpTo.TWO_DECIMALS)
                                count_purchase += 1
                                delta = 6
                                for col_num in range(4):
                                    c = ws.cell(
                                        row=row_num + 1, column=delta + col_num + 1
                                    )
                                    c.border = BORDER_BOTTOM_THIN
                                row_num += 1

                            purchase = next_purchase(purchases)
                    c = ws.cell(row=row_num - 1 + 1, column=10 + 1)
                    if len(purchases_price_formula) > 0:
                        c.value = "={}".format("+".join(purchases_price_formula))
                    else:
                        c.value = DECIMAL_ZERO
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_RED
                    ws.conditional_formatting.add(
                        f"{get_column_letter(11)}{row_num}",
                        CellIsRule(
                            operator="notEqual",
                            formula=[str(purchases_price)],
                            fill=yellow_fill,
                        ),
                    )
                if at_least_one_product:
                    row_num -= 1
            else:
                for offer_item in OfferItemReadOnly.objects.filter(
                    permanence_id=permanence.id,
                    producer_id=producer.id,
                ).order_by("preparation_sort_order_v2"):
                    purchase_set = (
                        Purchase.objects.filter(offer_item_id=offer_item.id)
                        .order_by(
                            "quantity_for_preparation_sort_order",
                            "customer__short_basket_name",
                        )
                        .select_related(
                            "customer",
                            "offer_item",
                            "offer_item__department_for_customer",
                        )
                    )
                    if delivery_id is not None:
                        purchase_set = purchase_set.filter(
                            customer_invoice__delivery_id=delivery_id
                        )
                    purchases = purchase_set.iterator()
                    purchase = next_purchase(purchases)
                    while purchase is not None:
                        at_least_one_product = True
                        while purchase is not None:
                            department_for_customer_save = (
                                purchase.offer_item.department_for_customer
                            )
                            department_for_customer_save__short_name = (
                                department_for_customer_save.short_name_v2
                                if department_for_customer_save is not None
                                else None
                            )
                            offer_item_save = purchase.offer_item
                            row_start_offer_item = row_num + 2
                            count_offer_item = 0
                            for col_num in range(11):
                                c = ws.cell(row=row_num + 1, column=col_num + 1)
                                c.border = BORDER_BOTTOM_THIN
                            purchases_quantity = DECIMAL_ZERO
                            row_num += 1
                            while (
                                purchase is not None
                                and offer_item_save == purchase.offer_item
                            ):
                                qty = purchase.get_producer_quantity()
                                if qty != DECIMAL_ZERO:
                                    base_unit = get_base_unit(
                                        offer_item_save.order_unit, qty
                                    )
                                    c = ws.cell(row=row_num + 1, column=0 + 1)
                                    c.value = purchase.id
                                    c = ws.cell(row=row_num + 1, column=1 + 1)
                                    c.value = purchase.offer_item_id
                                    if count_offer_item == 0:
                                        c = ws.cell(row=row_num + 1, column=2 + 1)
                                        c.value = "{}".format(
                                            purchase.offer_item.get_placement_display()
                                        )
                                        if placement_save is None:
                                            placement_save = c.value
                                        elif hide_column_placement:
                                            if placement_save != c.value:
                                                hide_column_placement = False
                                        c = ws.cell(row=row_num + 1, column=3 + 1)
                                        c.value = "{}".format(
                                            producer_save.short_profile_name
                                        )
                                    c = ws.cell(row=row_num + 1, column=5 + 1)
                                    if (
                                        department_for_customer_save__short_name
                                        is not None
                                    ):
                                        c.value = "{} - {}".format(
                                            purchase.get_long_name_with_customer_price(),
                                            department_for_customer_save__short_name,
                                        )
                                    else:
                                        c.value = "{}".format(
                                            purchase.get_long_name_with_customer_price()
                                        )
                                    c.alignment = ALIGN_LEFT_WRAP
                                    if count_offer_item != 0:
                                        c.font = FONT_GRAY
                                    c.number_format = FORMAT_TEXT
                                    c = ws.cell(row=row_num + 1, column=6 + 1)
                                    c.value = "{} - {}".format(
                                        purchase.customer.preparation_order,
                                        purchase.customer.short_basket_name,
                                    )
                                    c.number_format = FORMAT_TEXT
                                    c = ws.cell(row=row_num + 1, column=7 + 1)
                                    c.value = qty
                                    c.number_format = "#,##0.????"
                                    c.font = FONT_RED
                                    ws.conditional_formatting.add(
                                        f"{get_column_letter(8)}{row_num + 1}",
                                        CellIsRule(
                                            operator="notEqual",
                                            formula=[str(qty)],
                                            fill=yellow_fill,
                                        ),
                                    )
                                    c = ws.cell(row=row_num + 1, column=8 + 1)
                                    c.value = "{}".format(base_unit)
                                    c.number_format = FORMAT_TEXT
                                    if offer_item_save.order_unit == OrderUnit.PC_KG:
                                        c = ws.cell(row=row_num + 1, column=9 + 1)
                                        if offer_item_save.wrapped:
                                            c.value = "{} :".format(
                                                repanier.apps.REPANIER_SETTINGS_CURRENCY_DISPLAY
                                            )
                                        else:
                                            c.value = "{}".format(_("kg :"))
                                        c.number_format = FORMAT_TEXT
                                    else:
                                        if offer_item_save.wrapped:
                                            c = ws.cell(row=row_num + 1, column=9 + 1)
                                            c.value = "{} :".format(
                                                repanier.apps.REPANIER_SETTINGS_CURRENCY_DISPLAY
                                            )
                                            c.number_format = FORMAT_TEXT
                                    purchases_quantity += qty
                                    count_offer_item += 1
                                    delta = 6
                                    for col_num in range(4):
                                        c = ws.cell(
                                            row=row_num + 1, column=delta + col_num + 1
                                        )
                                        c.border = BORDER_BOTTOM_THIN
                                    row_num += 1

                                purchase = next_purchase(purchases)
                            if count_offer_item > 1:
                                c = ws.cell(row=row_num - 1 + 1, column=10 + 1)
                                c.value = "=SUM(H{}:H{})".format(
                                    row_start_offer_item, row_num
                                )
                                c.number_format = "#,##0.????"
                                if (
                                    not offer_item_save.wrapped
                                    and offer_item_save.order_unit
                                    in [OrderUnit.KG, OrderUnit.PC_KG]
                                ):
                                    c.font = FONT_RED
                                ws.conditional_formatting.add(
                                    f"{get_column_letter(11)}{row_num}",
                                    CellIsRule(
                                        operator="notEqual",
                                        formula=[str(purchases_quantity)],
                                        fill=yellow_fill,
                                    ),
                                )
                            row_num -= 1

            if at_least_one_product:
                for col_num in range(11):
                    c = ws.cell(row=row_num + 1, column=col_num + 1)
                    c.border = BORDER_BOTTOM_MEDIUM_DASHED
                row_num += 1
            producer = next_row(producers)
        ws.column_dimensions[get_column_letter(1)].hidden = True
        ws.column_dimensions[get_column_letter(2)].hidden = True
        if hide_column_placement:
            ws.column_dimensions[get_column_letter(3)].hidden = True
        if producer_counter <= 1:
            ws.column_dimensions[get_column_letter(4)].hidden = True
        ws.column_dimensions[get_column_letter(5)].hidden = True
    return wb


def export_customer(
    permanence=None,
    customer_invoice=None,
    deliveries_id=(),
    deposit=False,
    xlsx_formula=True,
    wb=None,
    ws_preparation_title=None,
):
    """Exporte le contrôle du panier client ou les consignes."""
    header = [
        (_("Purchase"), 5),
        (_("OfferItem"), 5),
        (_("Placement"), 15),
        (_("Producer"), 15),
        (_("Product"), 60),
        (_("Quantity"), 10),
        (_("Unit"), 10),
        (_("Unit price"), 10),
        (_("Deposit"), 10),
        (_("Total Price"), 12),
        (_("Basket"), 20),
    ]
    if len(deliveries_id) == 0:
        return export_customer_for_a_delivery(
            customer_invoice,
            0,
            None,
            deposit,
            header,
            permanence,
            wb,
            ws_preparation_title,
            FILL_YELLOW,
            xlsx_formula,
        )
    else:
        for delivery_cpt, delivery_id in enumerate(deliveries_id):
            wb = export_customer_for_a_delivery(
                customer_invoice,
                delivery_cpt,
                delivery_id,
                deposit,
                header,
                permanence,
                wb,
                ws_preparation_title,
                FILL_YELLOW,
                xlsx_formula,
            )
        return wb


def export_customer_for_a_delivery(
    customer_invoice,
    delivery_cpt,
    delivery_id,
    deposit,
    header,
    permanence,
    wb,
    ws_preparation_title,
    yellow_fill,
    xlsx_formula,
):
    """Génère le contrôle panier pour un point de livraison."""
    from repanier.apps import REPANIER_SETTINGS_CONFIG

    language_code = translation.get_language()
    if customer_invoice is not None:
        customer = customer_invoice.customer
        translation.activate(customer.language)
        if deposit:
            purchase_set = (
                Purchase.objects.filter(
                    permanence_id=permanence.id,
                    customer_id=customer_invoice.customer_id,
                    producer__isnull=False,
                    offer_item__order_unit=OrderUnit.DEPOSIT,
                )
                .order_by(
                    "offer_item__long_name_v2",
                    "offer_item__order_average_weight",
                )
                .select_related(
                    "customer", "offer_item", "offer_item__department_for_customer"
                )
            )
        else:
            purchase_set = (
                Purchase.objects.filter(
                    permanence_id=permanence.id,
                    customer_id=customer_invoice.customer_id,
                    producer__isnull=False,
                )
                .exclude(offer_item__order_unit=OrderUnit.DEPOSIT)
                .order_by(
                    "offer_item__long_name_v2",
                    "offer_item__order_average_weight",
                )
                .select_related(
                    "customer", "offer_item", "offer_item__department_for_customer"
                )
            )
    else:
        if deposit:
            purchase_set = (
                Purchase.objects.filter(
                    permanence_id=permanence.id,
                    producer__isnull=False,
                    offer_item__order_unit=OrderUnit.DEPOSIT,
                )
                .order_by(
                    "customer__short_basket_name",
                    "offer_item__producer",
                    "offer_item__long_name_v2",
                    "offer_item__order_average_weight",
                )
                .select_related(
                    "customer", "offer_item", "offer_item__department_for_customer"
                )
            )
        else:
            purchase_set = (
                Purchase.objects.filter(
                    permanence_id=permanence.id,
                    producer__isnull=False,
                )
                .exclude(offer_item__order_unit=OrderUnit.DEPOSIT)
                .order_by(
                    "customer__short_basket_name",
                    "offer_item__producer",
                    "offer_item__long_name_v2",
                    "offer_item__order_average_weight",
                )
                .select_related(
                    "customer", "offer_item", "offer_item__department_for_customer"
                )
            )
        if delivery_id is not None:
            purchase_set = purchase_set.filter(
                customer_invoice__delivery_id=delivery_id
            )
    purchases = purchase_set.iterator()
    purchase = next_purchase(purchases)
    if purchase is not None:
        config = REPANIER_SETTINGS_CONFIG
        group_label = config.group_label_v2
        if deposit:
            wb, ws = new_portrait_a4_sheet(
                wb,
                _("Deposits")
                if delivery_id is None
                else "{}-{}".format(delivery_cpt, _("Deposits")),
                permanence,
                header,
            )
        else:
            wb, ws = new_landscape_a4_sheet(
                wb,
                _("Basket control")
                if delivery_id is None
                else "{}-{}".format(delivery_cpt, _("Basket control")),
                permanence,
                header,
            )
        hide_column_placement = True
        hide_column_producer = True
        offer_item_save = purchase.offer_item
        placement_save = offer_item_save.placement
        producer_save = offer_item_save.producer
        row_num = 1
        if delivery_id is not None:
            c = ws.cell(row=row_num + 1, column=4 + 1)
            c.font = FONT_BOLD
            c.value = "{}".format(
                DeliveryBoard.objects.filter(id=delivery_id)
                .first()
                .get_delivery_display()
            )
            row_num += 1
        if ws_preparation_title is not None and xlsx_formula:
            if delivery_id is None:
                ref_preparation_sheet = format_worksheet_title(ws_preparation_title)
            else:
                ref_preparation_sheet = "{}-{}".format(
                    delivery_cpt, ws_preparation_title
                )
        else:
            ref_preparation_sheet = None
        while purchase is not None:
            customer_save = purchase.customer
            row_start_customer = row_num + 1
            first_purchase = True
            total_price = DECIMAL_ZERO
            something_ordered = False
            while purchase is not None and customer_save.id == purchase.customer_id:
                department_for_customer_save__id = (
                    offer_item_save.department_for_customer_id
                )
                department_for_customer_save__short_name = (
                    offer_item_save.department_for_customer.short_name_v2
                    if offer_item_save.department_for_customer is not None
                    else None
                )
                while (
                    purchase is not None
                    and customer_save.id == purchase.customer_id
                    and offer_item_save.department_for_customer_id
                    == department_for_customer_save__id
                ):
                    if placement_save != offer_item_save.placement:
                        hide_column_placement = False
                    if producer_save != offer_item_save.producer:
                        hide_column_producer = False
                    qty = purchase.get_producer_quantity()
                    if deposit or qty != DECIMAL_ZERO:
                        something_ordered = True
                        base_unit = get_base_unit(offer_item_save.order_unit, qty)
                        c = ws.cell(row=row_num + 1, column=0 + 1)
                        c.value = purchase.id
                        c = ws.cell(row=row_num + 1, column=1 + 1)
                        c.value = offer_item_save.id
                        c = ws.cell(row=row_num + 1, column=2 + 1)
                        c.value = "{}".format(offer_item_save.get_placement_display())
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        c = ws.cell(row=row_num + 1, column=3 + 1)
                        c.value = "{}".format(purchase.producer.short_profile_name)
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        c = ws.cell(row=row_num + 1, column=4 + 1)
                        if department_for_customer_save__short_name is not None:
                            c.value = "{} - {}".format(
                                purchase.get_long_name_with_customer_price(),
                                department_for_customer_save__short_name,
                            )
                        else:
                            c.value = "{}".format(
                                purchase.get_long_name_with_customer_price()
                            )
                        c.number_format = FORMAT_TEXT
                        c.alignment = ALIGN_LEFT_WRAP
                        c.border = BORDER_BOTTOM_THIN
                        c = ws.cell(row=row_num + 1, column=5 + 1)
                        if ref_preparation_sheet is None:
                            c.value = qty
                        else:
                            c.value = "=SUMIF('{}'!A:A,A{},'{}'!H:H)".format(
                                ref_preparation_sheet,
                                row_num + 1,
                                ref_preparation_sheet,
                            )

                        c.number_format = "#,##0.???"
                        c.border = BORDER_BOTTOM_THIN
                        if not deposit:
                            c.font = FONT_RED
                        if xlsx_formula:
                            ws.conditional_formatting.add(
                                f"{get_column_letter(6)}{row_num + 1}",
                                CellIsRule(
                                    operator="notEqual",
                                    formula=[str(qty)],
                                    fill=yellow_fill,
                                ),
                            )
                        c = ws.cell(row=row_num + 1, column=6 + 1)
                        c.value = "{}".format(base_unit)
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        c = ws.cell(row=row_num + 1, column=7 + 1)
                        customer_unit_price = purchase.get_customer_unit_price()
                        c.value = customer_unit_price
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                        c.border = BORDER_BOTTOM_THIN
                        c = ws.cell(row=row_num + 1, column=8 + 1)
                        unit_deposit = offer_item_save.unit_deposit.amount
                        c.value = unit_deposit
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                        c.border = BORDER_BOTTOM_THIN
                        c = ws.cell(row=row_num + 1, column=9 + 1)
                        if offer_item_save.order_unit == OrderUnit.PC_KG:
                            price_qty = qty * offer_item_save.order_average_weight
                            c.value = "=ROUND(F{}*{}*(H{}+I{}),2)".format(
                                row_num + 1,
                                offer_item_save.order_average_weight,
                                row_num + 1,
                                row_num + 1,
                            )
                        else:
                            price_qty = qty
                            c.value = "=ROUND(F{}*(H{}+I{}),2)".format(
                                row_num + 1, row_num + 1, row_num + 1
                            )
                        purchases_price = (
                            price_qty * (customer_unit_price + unit_deposit)
                        ).quantize(RoundUpTo.TWO_DECIMALS)
                        if not xlsx_formula:
                            c.value = purchases_price
                            total_price += purchases_price
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                        c.border = BORDER_BOTTOM_THIN
                        if xlsx_formula:
                            ws.conditional_formatting.add(
                                f"{get_column_letter(10)}{row_num + 1}",
                                CellIsRule(
                                    operator="notEqual",
                                    formula=[str(purchases_price)],
                                    fill=yellow_fill,
                                ),
                            )

                        c = ws.cell(row=row_num + 1, column=10 + 1)
                        if xlsx_formula:
                            c.value = "{} - {}".format(
                                purchase.customer.preparation_order,
                                purchase.customer.short_basket_name,
                            )
                        else:
                            c.value = "{}".format(purchase.customer.long_basket_name)
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        if first_purchase:
                            c.font = FONT_BOLD_ITALIC
                            first_purchase = False
                        else:
                            c.font = FONT_BOLD
                        row_num += 1
                    purchase = next_purchase(purchases)
                    offer_item_save = (
                        purchase.offer_item if purchase is not None else None
                    )
            if something_ordered:
                ci = CustomerInvoice.objects.filter(
                    permanence_id=permanence.id, customer_id=customer_save.id
                ).first()
                delta_transport = ci.delta_transport.amount
                if delta_transport:
                    c = ws.cell(row=row_num + 1, column=4 + 1)
                    c.value = "{}".format(_("Shipping cost"))
                    c.number_format = FORMAT_TEXT
                    c = ws.cell(row=row_num + 1, column=9 + 1)
                    c.value = delta_transport
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                c = ws.cell(row=row_num + 1, column=8 + 1)
                c.value = "{}".format(permanence)
                c.number_format = FORMAT_TEXT
                c.font = FONT_ITALIC
                c.alignment = ALIGN_RIGHT
                row_num += 1
                c = ws.cell(row=row_num + 1, column=4 + 1)
                c.value = "{}".format(group_label)
                c.number_format = FORMAT_TEXT
                c.font = FONT_BOLD
                c = ws.cell(row=row_num + 1, column=8 + 1)
                c.value = "{} {}".format(
                    _("Total Price"), customer_save.long_basket_name
                )
                c.number_format = FORMAT_TEXT
                c.font = FONT_BOLD
                c.alignment = ALIGN_RIGHT
                c = ws.cell(row=row_num + 1, column=9 + 1)
                if xlsx_formula:
                    c.value = "=SUM(J{}:J{})".format(row_start_customer, row_num)
                else:
                    c.value = total_price
                c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                c.font = FONT_BOLD
                row_num += 1
                for col_num in range(11):
                    c = ws.cell(row=row_num + 1, column=col_num + 1)
                    c.border = BORDER_BOTTOM_MEDIUM_DASHED
                row_num += 2

        ws.column_dimensions[get_column_letter(1)].hidden = True
        ws.column_dimensions[get_column_letter(2)].hidden = True
        if hide_column_placement:
            ws.column_dimensions[get_column_letter(3)].hidden = True
        if hide_column_producer:
            ws.column_dimensions[get_column_letter(4)].hidden = True
    translation.activate(language_code)
    return wb