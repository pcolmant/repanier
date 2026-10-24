"""
Excel export utilities for Repanier offers.
Handles export of product offers for permanences in PLANNED or OPENED status.
"""

import logging
from typing import Optional

from django.utils.translation import gettext_lazy as _
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

import repanier.apps
from repanier.const import DECIMAL_ZERO, LIMIT_ORDER_QTY_ITEM, SaleStatus, EMPTY_STRING
from repanier.models.offeritem import OfferItem
from repanier.models.producer import Producer
from repanier.models.product import Product
from repanier.xlsx.export_tools import (
    new_landscape_a4_sheet,
    worksheet_set_header,
    get_column_letter,
    ROW_VALUE,
    ROW_FORMAT,
)

# Constants for Excel formatting in offer exports
OFFER_SEPARATOR_COLUMN_WIDTH = 2.3
OFFER_QUANTITY_COLUMN_WIDTH = 20
OFFER_SEPARATOR_VALUE = "---"

logger = logging.getLogger(__name__)


def export_offer(permanence: "Permanence", wb: Optional[Workbook] = None) -> Workbook:
    """
    Export the product offer for a permanence to an Excel worksheet.

    Exports either:
    - Products from active producers (if permanence status is PLANNED)
    - Active offer items (if permanence status is OPENED)

    Args:
        permanence: The Permanence instance to export the offer for.
        wb: Optional existing OpenPyXL workbook. If None, a new workbook is created.

    Returns:
        Workbook: The OpenPyXL workbook containing the offer data.

    Raises:
        ValueError: If permanence is None.
        TypeError: If wb is not a Workbook instance or None.
    """
    if permanence is None:
        raise ValueError("permanence is required and cannot be None")

    if wb is not None and not isinstance(wb, Workbook):
        raise TypeError("wb must be an OpenPyXL Workbook instance or None")

    try:
        wb, ws = new_landscape_a4_sheet(wb, permanence, permanence)
    except Exception as e:
        logger.error(f"Failed to create worksheet for permanence {permanence}: {e}")
        raise

    current_row = 1
    items = []

    if permanence.status == SaleStatus.PLANNED.value:
        producers_in_this_permanence = Producer.objects.filter(
            permanence=permanence, is_active=True
        )
        items = (
            Product.objects.prefetch_related("producer", "department_for_customer")
            .filter(
                producer__in=producers_in_this_permanence,
                is_into_offer=True,
            )
            .order_by(
                "producer__short_profile_name",
                "department_for_customer",
                "long_name_v2",
                "order_average_weight",
            )
        )

    elif permanence.status == SaleStatus.OPENED.value:
        items = (
            OfferItem.objects.prefetch_related("producer", "department_for_customer")
            .filter(
                permanence_id=permanence.id,
                is_active=True,
            )
            .order_by(
                "order_sort_order_v2",
            )
        )

    for item in items:
        current_row = export_offer_row(item, current_row, ws)

    return wb


def export_offer_row(item: Product | OfferItem, current_row: int, ws: Worksheet) -> int:
    """
    Export a single item (Product or OfferItem) as a row in the worksheet.

    Args:
        item: Product or OfferItem instance to export. Can be either:
            - Product: When exporting from PLANNED status (pre-offer phase)
            - OfferItem: When exporting from OPENED status (active offer phase)
        current_row: Current row index (1-based). If 1, headers are written.
        ws: The OpenPyXL worksheet to write to.

    Returns:
        int: The next current_row index.

    Details:
        The row includes:
        - Basic info (Producer, Department, Product name)
        - Pricing (Producer price, Customer price, Deposit)
        - Quantity options (from min to alert, with steps)
    """
    try:
        row = [
            (_("Producer"), 15, item.producer.short_profile_name, FORMAT_TEXT),
            (
                _("Department"),
                15,
                item.department_for_customer.short_name_v2
                if item.department_for_customer is not None
                else EMPTY_STRING,
                FORMAT_TEXT,
            ),
            (
                _("Product"),
                60,
                item.get_long_name_with_customer_price(),
                FORMAT_TEXT,
            ),
            (
                _("Producer unit price"),
                10,
                item.producer_unit_price.amount
                if item.producer_unit_price < item.customer_unit_price
                else DECIMAL_ZERO,
                repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
            ),
            (
                _("Customer unit price"),
                10,
                item.customer_unit_price.amount,
                repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
            ),
            (
                _("Deposit"),
                10,
                item.unit_deposit.amount,
                repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
            ),
        ]

        # Use 1-based indices for OpenPyXL (rows and columns start at 1)
        current_col = 1

        if current_row == 1:
            worksheet_set_header(ws, row)
            current_row += 1

        # Write base row data using direct iteration
        for col_data in row:
            c = ws.cell(row=current_row, column=current_col)
            c.value, c.number_format = col_data[ROW_VALUE], col_data[ROW_FORMAT]
            current_col += 1

        # Separator column
        c = ws.cell(row=current_row, column=current_col)
        c.value = OFFER_SEPARATOR_VALUE
        ws.column_dimensions[get_column_letter(current_col)].width = OFFER_SEPARATOR_COLUMN_WIDTH
        c.number_format = FORMAT_TEXT
        current_col += 1

        # Quantity columns
        q_min = item.customer_minimum_order_quantity
        q_alert = item.get_q_alert()
        q_step = item.customer_increment_order_quantity
        q_valid = q_min
        q_counter = 0

        while q_valid <= q_alert and q_counter <= LIMIT_ORDER_QTY_ITEM:
            q_counter += 1
            c = ws.cell(row=current_row, column=current_col)
            c.value = item.get_display(
                qty=q_valid,
                order_unit=item.order_unit,
                unit_price_amount=item.customer_unit_price.amount,
            )
            ws.column_dimensions[get_column_letter(current_col)].width = OFFER_QUANTITY_COLUMN_WIDTH
            c.number_format = FORMAT_TEXT
            current_col += 1

            if q_valid < q_step:
                q_valid = q_step
            else:
                q_valid = q_valid + q_step

        return current_row + 1

    except AttributeError as e:
        logger.error(f"Invalid item type in export_offer_row (expected Product/OfferItem): {e}")
        raise
    except Exception as e:
        logger.error(f"Unexpected error exporting row {current_row}: {e}")
        raise
