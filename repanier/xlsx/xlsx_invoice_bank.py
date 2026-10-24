"""
Export Excel du tableau de bord des mouvements bancaires pour une permanence.
"""

from django.db.models import Q
from django.utils.translation import gettext_lazy as _
from openpyxl.styles.numbers import FORMAT_TEXT

import repanier.apps
from repanier.const import DECIMAL_ZERO, EMPTY_STRING
from repanier.models import CustomerInvoice
from repanier.models.bankaccount import BankAccount
from repanier.models.customer import Customer
from repanier.models.invoice import ProducerInvoice
from repanier.models.producer import Producer
from repanier.xlsx.export_tools import new_landscape_a4_sheet, worksheet_set_header
from repanier.xlsx.const import ROW_VALUE, ROW_FORMAT


def export_bank(permanence, wb=None, sheet_name=EMPTY_STRING):
    """
    Exporte le détail des mouvements bancaires (soldes, entrées/sorties de caisse) pour une permanence.
    """

    wb, ws = new_landscape_a4_sheet(
        wb, "{} {}".format(_("Dashboard"), sheet_name), permanence
    )

    row_num = 0

    bank_account = BankAccount.objects.filter(
        permanence_id=permanence.id, customer__isnull=True, producer__isnull=True
    ).first()
    if bank_account is None:
        # Permanence not invoiced yet : Nothing to do
        return wb

    customer_set = Customer.objects.filter(
        Q(customerinvoice__isnull=False, is_anonymized=False)
        | Q(customerinvoice__permanence_id=permanence.id)
    ).distinct()
    for customer in customer_set:
        bank_amount_in = bank_amount_out = DECIMAL_ZERO
        prepared = DECIMAL_ZERO
        customer_invoice = CustomerInvoice.objects.filter(
            customer_id=customer.id, permanence_id=permanence.id
        ).first()
        if customer_invoice is not None:
            balance_before = customer_invoice.previous_balance.amount
            bank_amount_in = customer_invoice.bank_amount_in.amount
            bank_amount_out = customer_invoice.bank_amount_out.amount
            if customer_invoice.customer_id == customer_invoice.customer_charged_id:
                prepared = customer_invoice.get_total_price_with_tax().amount
            balance_after = customer_invoice.balance.amount
        else:
            last_customer_invoice = (
                CustomerInvoice.objects.filter(
                    customer_id=customer.id, invoice_sort_order__lte=bank_account.id
                )
                .order_by("-invoice_sort_order")
                .first()
            )
            if last_customer_invoice is not None:
                balance_before = last_customer_invoice.balance.amount
                balance_after = last_customer_invoice.balance.amount
            else:
                balance_before = customer.initial_balance.amount
                balance_after = customer.initial_balance.amount

        row = [
            (_("Name"), 40, customer.long_basket_name, FORMAT_TEXT),
            (_("Previous balance"), 15, balance_before, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Cash in"), 10, bank_amount_in, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Cash out"), 10, bank_amount_out, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Prepared"), 10, prepared, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Final balance"), 15, balance_after, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Name"), 20, customer.short_basket_name, FORMAT_TEXT),
        ]

        if row_num == 0:
            worksheet_set_header(ws, row)
            row_num += 1

        current_row = row_num + 1
        current_col = 1
        for col_data in row:
            c = ws.cell(row=current_row, column=current_col)
            c.value = col_data[ROW_VALUE]
            c.number_format = col_data[ROW_FORMAT]
            current_col += 1

        row_num += 1

    row_break = row_num
    row_num += 1

    producer_set = Producer.objects.filter(producerinvoice__isnull=False).distinct()
    for producer in producer_set:
        bank_amount_in = bank_amount_out = DECIMAL_ZERO
        prepared = DECIMAL_ZERO
        producer_invoice = ProducerInvoice.objects.filter(
            producer_id=producer.id, permanence_id=permanence.id
        ).first()
        if producer_invoice is not None:
            balance_before = -producer_invoice.previous_balance.amount
            bank_amount_in = producer_invoice.bank_amount_in.amount
            bank_amount_out = producer_invoice.bank_amount_out.amount
            prepared = producer_invoice.get_total_price_with_tax().amount
            balance_after = -producer_invoice.balance.amount
        else:
            last_producer_invoice = (
                ProducerInvoice.objects.filter(
                    producer_id=producer.id, invoice_sort_order__lte=bank_account.id
                )
                .order_by("-invoice_sort_order")
                .first()
            )
            if last_producer_invoice is not None:
                balance_before = -last_producer_invoice.balance.amount
                balance_after = -last_producer_invoice.balance.amount
            else:
                balance_before = -producer.initial_balance.amount
                balance_after = -producer.initial_balance.amount

        row = [
            (_("Name"), 40, producer.long_profile_name, FORMAT_TEXT),
            (_("Previous balance"), 15, balance_before, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Cash in"), 10, bank_amount_in, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Cash out"), 10, bank_amount_out, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Prepared"), 10, prepared, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Final balance"), 15, balance_after, repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX),
            (_("Name"), 20, producer.short_profile_name, FORMAT_TEXT),
        ]

        if row_num == 0:
            worksheet_set_header(ws, row)
            row_num += 1

        current_row = row_num + 1
        current_col = 1
        for col_data in row:
            c = ws.cell(row=current_row, column=current_col)
            c.value = col_data[ROW_VALUE]
            c.number_format = col_data[ROW_FORMAT]
            current_col += 1

        row_num += 1

    final_bank_amount = bank_account.bank_amount_in.amount - bank_account.bank_amount_out.amount
    bank_account = BankAccount.objects.filter(
        id__lt=bank_account.id, customer__isnull=True, producer__isnull=True
    ).order_by("-id").first()
    if bank_account is not None:
        initial_bank_amount = bank_account.bank_amount_in.amount - bank_account.bank_amount_out.amount
    else:
        initial_bank_amount = DECIMAL_ZERO

    current_row = row_num + 1
    c = ws.cell(row=current_row, column=2)
    c.value = initial_bank_amount
    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX

    row_num += 1
    current_row = row_num + 1
    c = ws.cell(row=current_row, column=5)
    formula = "B{}+SUM(C{}:C{})-SUM(D{}:D{})".format(current_row, 2, current_row - 1, 2, current_row - 1)
    c.value = "=" + formula
    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX

    row_num += 1
    current_row = row_num + 1
    c = ws.cell(row=current_row, column=5)
    formula = "SUM(F{}:F{})-SUM(F{}:F{})".format(2, row_break, row_break + 2, current_row - 2)
    c.value = "=" + formula
    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX

    row_num += 1
    current_row = row_num + 1
    c = ws.cell(row=current_row, column=5)
    c.value = final_bank_amount
    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX

    return wb
