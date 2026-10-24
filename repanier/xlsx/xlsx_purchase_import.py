"""
Importation des données d'achat à partir de fichiers Excel.
"""

from decimal import Decimal
from django.db import transaction
from django.utils.translation import gettext_lazy as _
from openpyxl.reader.excel import load_workbook

from repanier.const import DECIMAL_ZERO, EMPTY_STRING, RoundUpTo, SaleStatus, DECIMAL_ONE
from repanier.models.offeritem import OfferItemReadOnly
from repanier.models.purchase import Purchase
from repanier.tools import cap
from repanier.xlsx.export_tools import format_worksheet_title
from repanier.xlsx.import_tools import get_header, get_row, get_customer_2_id_dict, get_producer_2_id_dict


@transaction.atomic
def import_purchase_sheet(
    worksheet, permanence=None, customer_2_id_dict=None, producer_2_id_dict=None
):
    """
    Importe les données d'achat depuis une feuille Excel.
    """
    error = False
    error_msg = None
    import_counter = 0
    header = get_header(worksheet)

    if not header:
        return True, _("No valid header found")

    current_row = 2
    array_purchase = []
    rule_of_3_source = DECIMAL_ZERO
    row = get_row(worksheet, header, current_row)

    while row and not error:
        try:
            row_format = row[_("Format")]
            if row_format in ["A", "B", "C", "D"]:
                import_counter += 1

                if row[_("Id")] is None:
                    error = True
                    error_msg = _("Row %(row_num)d : No purchase id given.") % {"row_num": current_row}
                    break

                purchase_id = Decimal(row[_("Id")])
                purchase = Purchase.objects.filter(id=purchase_id).first()

                if purchase is None:
                    error = True
                    error_msg = _("Row %(row_num)d : No purchase corresponding to the given purchase id.") % {"row_num": current_row}
                    break

                if purchase.permanence_id != permanence.id:
                    error = True
                    error_msg = _("Row %(row_num)d : The given permanence doesn't own the given purchase id.") % {"row_num": current_row}
                    break

                offer_item = OfferItemReadOnly.objects.filter(id=purchase.offer_item_id).first()
                if offer_item is None:
                    error = True
                    error_msg = _("Row %(row_num)d : No offer_item corresponding to the given purchase id.") % {"row_num": current_row}
                    break

                producer_id = producer_2_id_dict.get(row[_("Producer")])
                if producer_id != purchase.producer_id:
                    error = True
                    error_msg = _("Row %(row_num)d : No valid producer.") % {"row_num": current_row}
                    break

                customer_name = str(row[_("Customer")])
                customer_id = customer_2_id_dict.get(customer_name)
                if customer_id != purchase.customer_id:
                    error = True
                    error_msg = _("Row %(row_num)d : No valid customer.") % {"row_num": current_row}
                    break

                comment = cap(row[_("Comment")] or EMPTY_STRING, 100)
                quantity_has_been_modified = False

                producer_row_price = row[_("Purchase price")]
                if producer_row_price is not None:
                    producer_row_price = Decimal(producer_row_price).quantize(RoundUpTo.TWO_DECIMALS)
                    if purchase.purchase_price.amount != producer_row_price:
                        quantity_has_been_modified = True
                        producer_unit_price = (
                            offer_item.producer_unit_price.amount
                            + offer_item.unit_deposit.amount
                        ).quantize(RoundUpTo.TWO_DECIMALS)

                        if producer_unit_price != DECIMAL_ZERO:
                            purchase.quantity_invoiced = (producer_row_price / producer_unit_price).quantize(RoundUpTo.FOUR_DECIMALS)
                        else:
                            purchase.quantity_invoiced = DECIMAL_ZERO

                if not quantity_has_been_modified:
                    quantity_invoiced = (
                        DECIMAL_ZERO
                        if row[_("Quantity accounted for")] is None
                        else Decimal(row[_("Quantity accounted for")]).quantize(RoundUpTo.FOUR_DECIMALS)
                    )
                    if purchase.quantity_invoiced != quantity_invoiced:
                        purchase.quantity_invoiced = quantity_invoiced

                if row_format == "A":
                    array_purchase = []
                    rule_of_3_source = DECIMAL_ZERO
                    producer_unit_price = row[_("Producer unit price")]
                    if producer_unit_price is not None and not purchase.producer.invoice_by_basket:
                        previous_price = purchase.get_producer_unit_price()
                        if producer_unit_price != previous_price:
                            offer_item.producer_unit_price = producer_unit_price
                            offer_item.save()

                purchase.comment = comment
                purchase.save()
                rule_of_3_source += purchase.purchase_price.amount
                array_purchase.append(purchase)

            if row_format in ["C", "D"]:
                rule_of_3_target = row[_("Rule of 3")]
                if rule_of_3_target is not None:
                    rule_of_3_target = Decimal(rule_of_3_target).quantize(RoundUpTo.TWO_DECIMALS)
                    if rule_of_3_target != rule_of_3_source:
                        max_purchase_counter = len(array_purchase)
                        if max_purchase_counter <= 1:
                            error = True
                            error_msg = _("Row %(row_num)d : Rule of 3 target in wrong context.") % {"row_num": current_row}
                            return error, error_msg

                        ratio = rule_of_3_target / rule_of_3_source if rule_of_3_source != DECIMAL_ZERO else (DECIMAL_ZERO if rule_of_3_target == DECIMAL_ZERO else DECIMAL_ONE)

                        if ratio != DECIMAL_ONE:
                            adjusted_invoice = DECIMAL_ZERO
                            for i, purchase_item in enumerate(array_purchase, start=1):
                                producer_unit_price = purchase_item.offer_item.producer_unit_price.amount

                                if i == max_purchase_counter:
                                    delta = rule_of_3_target - adjusted_invoice
                                    purchase_item.quantity_invoiced = (delta / producer_unit_price).quantize(RoundUpTo.FOUR_DECIMALS) if producer_unit_price != DECIMAL_ZERO else DECIMAL_ZERO
                                else:
                                    purchase_item.quantity_invoiced = (purchase_item.quantity_invoiced * ratio).quantize(RoundUpTo.FOUR_DECIMALS)
                                    adjusted_invoice += (purchase_item.quantity_invoiced * producer_unit_price).quantize(RoundUpTo.TWO_DECIMALS)

                                purchase_item.save()

            current_row += 1
            row = get_row(worksheet, header, current_row)

        except KeyError as e:
            error = True
            error_msg = _("Row %(row_num)d : A required column is missing %(error_msg)s.") % {"row_num": current_row, "error_msg": str(e)}
        except Exception as e:
            error = True
            error_msg = _("Row %(row_num)d : %(error_msg)s.") % {"row_num": current_row, "error_msg": str(e)}

    if import_counter == 0:
        error = True
        error_msg = _("Nothing to import.")

    return error, error_msg


def handle_uploaded_purchase(request, permanence, file_to_import, *args):
    """
    Gère le fichier d'achat téléversé.
    """
    error = False
    error_msg = None

    wb = load_workbook(file_to_import)
    customer_buyinggroup_id, customer_2_id_dict = get_customer_2_id_dict()
    producer_buyinggroup_id, producer_2_id_dict = get_producer_2_id_dict()

    if customer_buyinggroup_id is None:
        error = True
        error_msg = _("At least one customer must represent the buying group.")
    elif producer_buyinggroup_id is None:
        error = True
        error_msg = _("At least one producer must represent the buying group.")

    if not error:
        if permanence is not None:
            if permanence.status == SaleStatus.SEND:
                ws_name = format_worksheet_title(permanence)
                if ws_name in wb.sheetnames:
                    ws = wb[ws_name]
                    error, error_msg = import_purchase_sheet(
                        ws,
                        permanence=permanence,
                        customer_2_id_dict=customer_2_id_dict,
                        producer_2_id_dict=producer_2_id_dict,
                    )
                    if error:
                        error_msg = f"{ws_name} > {error_msg}"
                else:
                    error = True
                    error_msg = _("The permanence doesn't exist.")
            else:
                error = True
                error_msg = _("The accounting has already been done")
        else:
            error = True
            error_msg = _("The permanence doesn't exist.")

    return error, error_msg