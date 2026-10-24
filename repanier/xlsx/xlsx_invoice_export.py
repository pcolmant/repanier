"""
Export Excel détaillé des factures et achats préparés.
"""

from django.utils.translation import gettext_lazy as _
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter

import repanier.apps
from repanier.const import (
    DECIMAL_ZERO,
    EMPTY_STRING,
    OrderUnit,
)
from repanier.models import CustomerInvoice
from repanier.models.purchase import Purchase
from repanier.tools import get_invoice_unit_display
from repanier.xlsx.export_tools import new_landscape_a4_sheet, worksheet_set_header
from repanier.xlsx.xlsx_styles import (
    FONT_BOLD,
    BORDER_THIN,
    BORDER_HAIR,
    BORDER_TOP_THIN_BOTTOM_THIN,
)
from repanier.xlsx.const import ROW_VALUE, ROW_FORMAT, ROW_BOX


def export_invoice(
    permanence=None,
    year=None,
    customer=None,
    producer=None,
    wb=None,
    sheet_name=EMPTY_STRING,
):
    """
    Exporte le détail des achats préparés pour facturation (filtrable par permanence, année, client ou producteur).
    """

    from repanier.apps import REPANIER_SETTINGS_CONFIG

    hide_producer_prices = producer is None
    hide_customer_prices = customer is None
    purchase_set = Purchase.objects.get_invoices(
        permanence=permanence, year=year, customer=customer, producer=producer
    )
    if purchase_set.exists():
        wb, ws = new_landscape_a4_sheet(wb, sheet_name, permanence or year)
        row = []
        row_num = 0
        hide_column_deposit = True
        permanence_id_save = None

        for purchase in purchase_set:

            if permanence_id_save is not None:
                row_num = shipping_cost_row(customer, permanence_id_save, row, row_num, ws)
                permanence_id_save = purchase.permanence_id

            qty = purchase.quantity_invoiced

            if purchase.offer_item.unit_deposit.amount != DECIMAL_ZERO:
                hide_column_deposit = False

            unit = get_invoice_unit_display(order_unit=purchase.offer_item.order_unit, qty=qty)
            row = [
                (
                    _("Permanence"),
                    20,
                    purchase.permanence.get_permanence_display(),
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Producer"),
                    15,
                    purchase.producer.short_profile_name,
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Basket"),
                    20,
                    purchase.customer.short_basket_name,
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Department"),
                    15,
                    purchase.offer_item.department_for_customer.short_name_v2
                    if purchase.offer_item.department_for_customer is not None
                    else EMPTY_STRING,
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Product"),
                    60,
                    purchase.get_long_name_with_customer_price(),
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Quantity"),
                    10,
                    qty,
                    "#,##0.????",
                    True
                    if purchase.offer_item.order_unit == OrderUnit.PC_KG
                    else False,
                ),
                (_("Unit"), 10, unit, FORMAT_TEXT, False),
                (
                    _("Deposit"),
                    10,
                    purchase.offer_item.unit_deposit.amount,
                    repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                    False,
                ),
            ]
            if hide_producer_prices:
                row += [
                    (_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False),
                    (_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False),
                    (_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False),
                ]
            else:
                row += [
                    (
                        _("Producer unit price"),
                        10,
                        purchase.get_producer_unit_price(),
                        repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                        False,
                    ),
                    (
                        _("Purchase price"),
                        10,
                        purchase.purchase_price.amount,
                        repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                        False,
                    ),
                    (
                        _("VAT"),
                        10,
                        purchase.producer_vat.amount,
                        repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                        False,
                    ),
                ]

            if hide_customer_prices:
                row += [
                    (_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False),
                    (_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False),
                    (_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False),
                ]
            else:
                row += [
                    (
                        _("Customer unit price"),
                        10,
                        purchase.get_customer_unit_price(),
                        repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                        False,
                    ),
                    (
                        _("Selling price"),
                        10,
                        purchase.selling_price.amount,
                        repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                        False,
                    ),
                    (
                        _("VAT"),
                        10,
                        purchase.customer_vat.amount,
                        repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX,
                        False,
                    ),
                ]
            if hide_producer_prices and hide_customer_prices:
                row += [(_("Empty"), 10, EMPTY_STRING, FORMAT_TEXT, False)]
            else:
                row += [
                    (
                        _("VAT rate"),
                        10,
                        purchase.offer_item.get_vat_level_display(),
                        FORMAT_TEXT,
                        False,
                    )
                ]
            row += [
                (_("Comment"), 30, purchase.comment, FORMAT_TEXT, False),
                (
                    _("Accounting entry status"),
                    10,
                    purchase.get_status_display(),
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Customer"),
                    10,
                    purchase.customer.user.email,
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Reference"),
                    10,
                    purchase.offer_item.reference,
                    FORMAT_TEXT,
                    False,
                ),
                (
                    _("Wrapped"),
                    10,
                    "{}".format(_("Yes") if purchase.offer_item.wrapped else _("No")),
                    FORMAT_TEXT,
                    False,
                ),
            ]

            if row_num == 0:
                worksheet_set_header(ws, row)
                row_num += 1

            # Use current_row/current_col to avoid repeated +1 calculations
            current_row = row_num + 1
            current_col = 1
            for col_num in range(len(row)):
                c = ws.cell(row=current_row, column=current_col)
                c.value = row[col_num][ROW_VALUE]
                c.number_format = row[col_num][ROW_FORMAT]
                if row[col_num][ROW_BOX]:
                    c.border = BORDER_THIN
                else:
                    c.border = BORDER_HAIR
                if col_num == 8:
                    c.font = FONT_BOLD
                current_col += 1

            row_num += 1

        if ws is not None:
            if hide_column_deposit:
                ws.column_dimensions[get_column_letter(8)].hidden = True
            if hide_producer_prices:
                ws.column_dimensions[get_column_letter(9)].hidden = True
                ws.column_dimensions[get_column_letter(10)].hidden = True
                ws.column_dimensions[get_column_letter(11)].hidden = True
            if hide_customer_prices:
                ws.column_dimensions[get_column_letter(12)].hidden = True
                ws.column_dimensions[get_column_letter(13)].hidden = True
                ws.column_dimensions[get_column_letter(14)].hidden = True
            if hide_producer_prices and hide_customer_prices:
                ws.column_dimensions[get_column_letter(15)].hidden = True
            ws.column_dimensions[get_column_letter(18)].hidden = True
            ws.column_dimensions[get_column_letter(19)].hidden = True
            ws.column_dimensions[get_column_letter(20)].hidden = True

            if permanence_id_save is not None:
                row_num = shipping_cost_row(customer, permanence_id_save, row, row_num, ws)

            # Use current_row/current_col for totals row
            current_row = row_num + 1
            current_col = 1
            for col_num in range(len(row)):
                c = ws.cell(row=current_row, column=current_col)
                c.border = BORDER_TOP_THIN_BOTTOM_THIN
                if col_num == 2:
                    c.value = "{} {}".format(_("Total Price"), sheet_name)
                    c.font = FONT_BOLD
                if col_num == 9:
                    formula = "SUM(J{}:J{})".format(2, row_num)
                    c.value = "=" + formula
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_BOLD
                if col_num == 10:
                    formula = "SUM(K{}:K{})".format(2, row_num)
                    c.value = "=" + formula
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_BOLD
                if col_num == 12:
                    formula = "SUM(M{}:M{})".format(2, row_num)
                    c.value = "=" + formula
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_BOLD
                if col_num == 13:
                    formula = "SUM(N{}:N{})".format(2, row_num)
                    c.value = "=" + formula
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_BOLD
                current_col += 1

            if customer is not None:
                config = REPANIER_SETTINGS_CONFIG
                group_label = config.group_label_v2
                if group_label:
                    row_num += 1
                    # Use current_row/current_col for group label row
                    current_row = row_num + 1
                    current_col = 1
                    for col_num in range(len(row)):
                        c = ws.cell(row=current_row, column=current_col)
                        c.border = BORDER_TOP_THIN_BOTTOM_THIN
                        if col_num == 2:
                            c.value = group_label
                            c.font = FONT_BOLD
                        current_col += 1
    return wb


def shipping_cost_row(customer, permanence_id_save, row, row_num, ws):
    """
    Add a row for shipping costs if applicable.
    """

    if customer is not None and permanence_id_save is not None:
        customer_invoice = CustomerInvoice.objects.filter(
            permanence_id=permanence_id_save, customer_id=customer.id
        ).first()
        delta_transport = customer_invoice.delta_transport.amount
        if delta_transport:
            delivery = customer_invoice.delivery or EMPTY_STRING
            # Use current_row/current_col for consistency
            current_row = row_num + 1
            current_col = 1
            for col_num in range(len(row)):
                c = ws.cell(row=current_row, column=current_col)
                c.border = BORDER_TOP_THIN_BOTTOM_THIN
                if col_num == 2:
                    c.value = "{} {}".format(_("Shipping cost"), delivery)
                    c.font = FONT_BOLD
                if col_num == 12:
                    c.value = delta_transport
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                current_col += 1
            row_num += 1
    return row_num
