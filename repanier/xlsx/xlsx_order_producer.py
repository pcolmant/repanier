"""
Producer-related Excel exports generation.
Uses the centralized style repository (xlsx_styles.py) while strictly maintaining
the original business logic and structure.
"""

from django.utils import translation
from django.utils.translation import gettext_lazy as _
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles.numbers import FORMAT_TEXT
from openpyxl.utils import get_column_letter

import repanier.apps
from repanier.const import *
from repanier.models.invoice import ProducerInvoice
from repanier.models.offeritem import OfferItemReadOnly
from repanier.models.purchase import Purchase
from repanier.tools import get_base_unit, next_row
from repanier.xlsx.export_tools import new_landscape_a4_sheet
from repanier.xlsx.xlsx_styles import (
    FILL_YELLOW,
    BORDER_BOTTOM_THIN,
    FONT_BOLD,
    FONT_RED,
    FONT_GRAY, ALIGN_LEFT_WRAP,
)


def export_producer_by_product(permanence, producer, wb=None):
    """Export a producer's order grouped by product."""
    language_code = translation.get_language()
    translation.activate(producer.language)

    header = [
        (_("Basket"), 20),
        (_("Quantity"), 10),
        (_("Unit"), 10),
        (_("Reference"), 20),
        (_("Product"), 60),
        (_("Unit price"), 10),
        (_("Deposit"), 10),
        (_("Total Price"), 12),
    ]
    show_column_reference = False
    hide_column_short_basket_name = True
    hide_column_unit_deposit = True
    formula_main_total = []
    offer_items = (
        OfferItemReadOnly.objects.filter(
            permanence_id=permanence.id,
            producer_id=producer.id,
            quantity_invoiced__gt=DECIMAL_ZERO,
        )
        .order_by("department_for_customer", "producer_sort_order_v2")
        .iterator()
    )
    offer_item = next_row(offer_items)
    if offer_item:
        wb, ws = new_landscape_a4_sheet(
            wb,
            "{} {}".format(producer.short_profile_name, _("by product")),
            permanence,
            header,
        )
        row_num = 1
        producer_purchases_price = DECIMAL_ZERO
        status = (
            ProducerInvoice.objects.filter(
                permanence_id=permanence.id, producer_id=producer.id
            )
            .only("status")
            .first()
            .status
        )
        while offer_item is not None:
            department_for_customer_save = offer_item.department_for_customer
            current_row = row_num + 1
            c = ws.cell(row=current_row, column=2)
            c.value = (
                "{}".format(department_for_customer_save.short_name_v2)
                if department_for_customer_save is not None
                else "---"
            )
            c.font = FONT_BOLD
            row_num += 1
            row_start_department = row_num
            department_purchases_price = DECIMAL_ZERO
            while (
                offer_item is not None
                and department_for_customer_save == offer_item.department_for_customer
            ):
                if offer_item.unit_deposit != DECIMAL_ZERO:
                    hide_column_unit_deposit = False
                show_column_reference = (
                    show_column_reference | len(offer_item.reference) < 36
                )
                if offer_item.wrapped:
                    hide_column_short_basket_name = False
                    first_purchase = True
                    purchase_set = Purchase.objects.filter(
                        offer_item_id=offer_item.id,
                    ).order_by(
                        "offer_item__producer_sort_order_v2",
                        "customer__short_basket_name",
                    )
                    for purchase in purchase_set:
                        if offer_item.stock > DECIMAL_ZERO:
                            qty = DECIMAL_ZERO
                        else:
                            qty = purchase.get_producer_quantity()
                        if qty != DECIMAL_ZERO:
                            base_unit = get_base_unit(offer_item.order_unit, qty)
                            current_row = row_num + 1
                            current_col = 1
                            c = ws.cell(row=current_row, column=current_col)
                            c.value = "{}".format(purchase.customer.short_basket_name)
                            c.number_format = FORMAT_TEXT
                            c.border = BORDER_BOTTOM_THIN
                            current_col += 1
                            c = ws.cell(row=current_row, column=current_col)
                            c.value = qty
                            c.number_format = "#,##0.???"
                            c.border = BORDER_BOTTOM_THIN
                            c.font = FONT_RED
                            ws.conditional_formatting.add(
                                f"{get_column_letter(2)}{current_row}",
                                CellIsRule(
                                    operator="notEqual",
                                    formula=[str(qty)],
                                    fill=FILL_YELLOW,
                                ),
                            )
                            current_col += 1
                            c = ws.cell(row=current_row, column=current_col)
                            c.value = "{}".format(base_unit)
                            c.number_format = FORMAT_TEXT
                            c.border = BORDER_BOTTOM_THIN
                            current_col += 1
                            c = ws.cell(row=current_row, column=current_col)
                            reference = (
                                offer_item.reference
                                if len(offer_item.reference) < 36
                                else EMPTY_STRING
                            )
                            if not reference.isdigit():
                                c.value = "[{}]".format(reference)
                            else:
                                c.value = "{}".format(reference)
                            c.number_format = FORMAT_TEXT
                            c.border = BORDER_BOTTOM_THIN
                            c = ws.cell(row=row_num, column=5)
                            c.value = offer_item.get_long_name_with_producer_price()
                            c.number_format = FORMAT_TEXT
                            c.alignment = ALIGN_LEFT_WRAP
                            c.border = BORDER_BOTTOM_THIN
                            if first_purchase:
                                first_purchase = False
                            else:
                                c.font = FONT_GRAY
                            current_col = 6
                            c = ws.cell(row=current_row, column=current_col)
                            producer_unit_price = purchase.get_producer_unit_price()
                            customer_unit_price = purchase.get_customer_unit_price()
                            if producer_unit_price < customer_unit_price:
                                unit_price = producer_unit_price
                            else:
                                unit_price = customer_unit_price
                            c.value = unit_price
                            c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                            c.border = BORDER_BOTTOM_THIN
                            current_col += 1
                            c = ws.cell(row=current_row, column=current_col)
                            unit_deposit = offer_item.unit_deposit.amount
                            c.value = unit_deposit
                            c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                            c.border = BORDER_BOTTOM_THIN
                            current_col += 1
                            c = ws.cell(row=current_row, column=current_col)
                            if offer_item.order_unit == OrderUnit.PC_KG:
                                price_qty = qty * offer_item.order_average_weight
                                c.value = "=ROUND(B{}*{}*(F{}+G{}),2)".format(
                                    current_row,
                                    offer_item.order_average_weight,
                                    current_row,
                                    current_row,
                                )
                            else:
                                price_qty = qty
                                c.value = "=ROUND(B{}*(F{}+G{}),2)".format(
                                    current_row, current_row, current_row
                                )
                            purchase_price = (
                                price_qty * (unit_price + unit_deposit)
                            ).quantize(RoundUpTo.TWO_DECIMALS)
                            c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                            c.border = BORDER_BOTTOM_THIN
                            department_purchases_price += purchase_price
                            ws.conditional_formatting.add(
                                f"{get_column_letter(8)}{current_row}",
                                CellIsRule(
                                    operator="notEqual",
                                    formula=[str(purchase_price)],
                                    fill=FILL_YELLOW,
                                ),
                            )
                            row_num += 1
                else:
                    qty = offer_item.get_producer_quantity(status)
                    if qty != DECIMAL_ZERO:
                        base_unit = get_base_unit(offer_item.order_unit, qty)
                        current_row = row_num + 1
                        current_col = 1
                        c = ws.cell(row=current_row, column=current_col)
                        c.value = "{}".format(settings.REPANIER_SETTINGS_GROUP_NAME)
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        c.value = qty
                        c.number_format = "#,##0.???"
                        c.border = BORDER_BOTTOM_THIN
                        c.font = FONT_RED
                        ws.conditional_formatting.add(
                            f"{get_column_letter(2)}{current_row}",
                            CellIsRule(
                                operator="notEqual",
                                formula=[str(qty)],
                                fill=FILL_YELLOW,
                            ),
                        )
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        c.value = "{}".format(base_unit)
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        reference = (
                            offer_item.reference
                            if len(offer_item.reference) < 36
                            else EMPTY_STRING
                        )
                        if not reference.isdigit():
                            c.value = "[{}]".format(reference)
                        else:
                            c.value = "{}".format(reference)
                        c.number_format = FORMAT_TEXT
                        c.border = BORDER_BOTTOM_THIN
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        c.value = offer_item.get_long_name_with_producer_price()
                        c.number_format = FORMAT_TEXT
                        c.alignment = ALIGN_LEFT_WRAP
                        c.border = BORDER_BOTTOM_THIN
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        producer_unit_price = offer_item.producer_unit_price.amount
                        customer_unit_price = offer_item.customer_unit_price.amount
                        if producer_unit_price < customer_unit_price:
                            unit_price = producer_unit_price
                        else:
                            unit_price = customer_unit_price
                        c.value = unit_price
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                        c.border = BORDER_BOTTOM_THIN
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        unit_deposit = offer_item.unit_deposit.amount
                        c.value = unit_deposit
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                        c.border = BORDER_BOTTOM_THIN
                        current_col += 1
                        c = ws.cell(row=current_row, column=current_col)
                        if offer_item.order_unit == OrderUnit.PC_KG:
                            price_qty = qty * offer_item.order_average_weight
                            c.value = "=ROUND(B{}*{}*(F{}+G{}),2)".format(
                                current_row,
                                offer_item.order_average_weight,
                                current_row,
                                current_row,
                            )
                        else:
                            price_qty = qty
                            c.value = "=ROUND(B{}*(F{}+G{}),2)".format(
                                current_row, current_row, current_row
                            )
                        purchase_price = (
                            price_qty * (unit_price + unit_deposit)
                        ).quantize(RoundUpTo.TWO_DECIMALS)
                        c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                        c.border = BORDER_BOTTOM_THIN
                        department_purchases_price += purchase_price
                        ws.conditional_formatting.add(
                            f"{get_column_letter(8)}{current_row}",
                            CellIsRule(
                                operator="notEqual",
                                formula=[str(purchase_price)],
                                fill=FILL_YELLOW,
                            ),
                        )

                        row_num += 1
                offer_item = next_row(offer_items)
            current_row = row_num + 1
            current_col = 1
            for col_num in range(8):
                c = ws.cell(row=current_row, column=current_col)
                c.border = BORDER_BOTTOM_THIN
                if col_num == 4:
                    if producer.producer_price_are_wo_vat:
                        c.value = "{} {} {}".format(
                            _("Total Price"),
                            _("wo VAT"),
                            department_for_customer_save.short_name_v2
                            if department_for_customer_save is not None
                            else "---",
                        )
                    else:
                        c.value = "{} {} {}".format(
                            _("Total Price"),
                            _("w TVA"),
                            department_for_customer_save.short_name_v2
                            if department_for_customer_save is not None
                            else "---",
                        )
                    c.number_format = FORMAT_TEXT
                if col_num == 7:
                    formula = "SUM(H{}:H{})".format(row_start_department, row_num)
                    c.value = "=" + formula
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_BOLD
                    producer_purchases_price += department_purchases_price
                    ws.conditional_formatting.add(
                        f"{get_column_letter(8)}{current_row}",
                        CellIsRule(
                            operator="notEqual",
                            formula=[str(department_purchases_price)],
                            fill=FILL_YELLOW,
                        ),
                    )
                    formula_main_total.append(formula)
                current_col += 1
            row_num += 1
        current_row = row_num + 1
        current_col = 1
        for col_num in range(8):
            c = ws.cell(row=current_row, column=current_col)
            c.border = BORDER_BOTTOM_THIN
            if col_num == 1:
                if producer.producer_price_are_wo_vat:
                    c.value = "{} {} {}".format(
                        _("Total Price"),
                        _("wo VAT"),
                        settings.REPANIER_SETTINGS_GROUP_NAME,
                    )
                else:
                    c.value = "{} {} {}".format(
                        _("Total Price"),
                        _("w TVA"),
                        settings.REPANIER_SETTINGS_GROUP_NAME,
                    )
                c.number_format = FORMAT_TEXT
            if col_num == 7:
                c.value = "=" + "+".join(formula_main_total)
                c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                c.font = FONT_BOLD
                ws.conditional_formatting.add(
                    f"{get_column_letter(8)}{current_row}",
                    CellIsRule(
                        operator="notEqual",
                        formula=[str(producer_purchases_price)],
                        fill=FILL_YELLOW,
                    ),
                )
            current_col += 1

        if hide_column_short_basket_name:
            ws.column_dimensions[get_column_letter(1)].hidden = True
        if not show_column_reference:
            ws.column_dimensions[get_column_letter(4)].hidden = True
        if hide_column_unit_deposit:
            ws.column_dimensions[get_column_letter(7)].hidden = True
    translation.activate(language_code)
    return wb


def export_producer_by_customer(permanence, producer, wb=None):
    """Export a producer's order grouped by customer basket."""
    language_code = translation.get_language()
    translation.activate(producer.language)

    header = [
        (_("Quantity"), 10),
        (_("Unit"), 10),
        (_("Reference"), 20),
        (_("Product"), 60),
        (_("Unit price"), 10),
        (_("Deposit"), 10),
        (_("Total Price"), 12),
        (EMPTY_STRING, 15),
    ]
    show_column_reference = False
    hide_column_unit_deposit = True
    formula_main_total = []
    purchase_set = (
        Purchase.objects.filter(
            permanence_id=permanence.id,
            producer_id=producer.id,
        )
        .order_by("customer__short_basket_name", "offer_item__producer_sort_order_v2")
        .select_related("offer_item")
    )
    purchases = purchase_set.iterator()
    purchase = next_row(purchases)
    if purchase:
        wb, ws = new_landscape_a4_sheet(
            wb,
            "{} {}".format(producer.short_profile_name, _("duplicate, by basket")),
            permanence,
            header,
        )
        row_num = 1
        while purchase is not None:
            customer_save = purchase.customer
            current_row = row_num + 1
            c = ws.cell(row=current_row, column=1)
            if customer_save.zero_waste:
                c.value = "{} --------------- {}".format(
                    customer_save.short_basket_name, _("Zero waste")
                )
                c.fill = FILL_YELLOW
            else:
                c.value = "{}".format(customer_save.short_basket_name)
            c.font = FONT_BOLD
            row_num += 1
            row_start_customer = row_num
            first_purchase = True
            while purchase is not None and customer_save.id == purchase.customer_id:
                qty = purchase.get_producer_quantity()
                if qty != DECIMAL_ZERO:
                    offer_item_save = purchase.offer_item
                    base_unit = get_base_unit(purchase.offer_item.order_unit, qty)
                    current_row = row_num + 1
                    current_col = 1
                    c = ws.cell(row=current_row, column=current_col)
                    c.value = qty
                    c.number_format = "#,##0.???"
                    c.border = BORDER_BOTTOM_THIN
                    c.font = FONT_BOLD
                    current_col += 1
                    c = ws.cell(row=current_row, column=current_col)
                    c.value = "{}".format(base_unit)
                    c.number_format = FORMAT_TEXT
                    c.border = BORDER_BOTTOM_THIN
                    current_col += 1
                    c = ws.cell(row=current_row, column=current_col)
                    reference = (
                        offer_item_save.reference
                        if len(offer_item_save.reference) < 36
                        else EMPTY_STRING
                    )
                    if not reference.isdigit():
                        c.value = "[{}]".format(reference)
                    else:
                        c.value = "{}".format(reference)
                    c.number_format = FORMAT_TEXT
                    c.border = BORDER_BOTTOM_THIN
                    current_col += 1
                    c = ws.cell(row=current_row, column=current_col)
                    c.value = offer_item_save.get_long_name_with_producer_price()
                    c.number_format = FORMAT_TEXT
                    c.alignment = ALIGN_LEFT_WRAP
                    c.border = BORDER_BOTTOM_THIN
                    if first_purchase:
                        first_purchase = False
                    current_col += 1
                    c = ws.cell(row=current_row, column=current_col)
                    producer_unit_price = purchase.get_producer_unit_price()
                    customer_unit_price = purchase.get_customer_unit_price()
                    if producer_unit_price < customer_unit_price:
                        unit_price = producer_unit_price
                    else:
                        unit_price = customer_unit_price
                    c.value = unit_price
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.border = BORDER_BOTTOM_THIN
                    current_col += 1
                    c = ws.cell(row=current_row, column=current_col)
                    c.value = offer_item_save.unit_deposit.amount
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.border = BORDER_BOTTOM_THIN
                    current_col += 1
                    c = ws.cell(row=current_row, column=current_col)
                    if offer_item_save.order_unit == OrderUnit.PC_KG:
                        c.value = "=ROUND(A{}*{}*(E{}+F{}),2)".format(
                            current_row,
                            offer_item_save.order_average_weight,
                            current_row,
                            current_row,
                        )
                    else:
                        c.value = "=ROUND(A{}*(E{}+F{}),2)".format(
                            current_row, current_row, current_row
                        )
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.border = BORDER_BOTTOM_THIN
                    row_num += 1
                purchase = next_row(purchases)
            current_row = row_num + 1
            current_col = 1
            for col_num in range(7):
                c = ws.cell(row=current_row, column=current_col)
                c.border = BORDER_BOTTOM_THIN
                if col_num == 3:
                    if producer.producer_price_are_wo_vat:
                        c.value = "{} {} {}".format(
                            _("Total Price"),
                            _("wo VAT"),
                            customer_save.short_basket_name,
                        )
                    else:
                        c.value = "{} {} {}".format(
                            _("Total Price"),
                            _("w TVA"),
                            customer_save.short_basket_name,
                        )
                    c.number_format = FORMAT_TEXT
                if col_num == 6:
                    formula = "SUM(G{}:G{})".format(row_start_customer, row_num)
                    c.value = "=" + formula
                    c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                    c.font = FONT_BOLD
                    formula_main_total.append(formula)
                current_col += 1
            row_num += 1
        current_row = row_num + 1
        current_col = 1
        for col_num in range(7):
            c = ws.cell(row=current_row, column=current_col)
            c.border = BORDER_BOTTOM_THIN
            if col_num == 0:
                if producer.producer_price_are_wo_vat:
                    c.value = "{} {} {}".format(
                        _("Total Price"),
                        _("wo VAT"),
                        settings.REPANIER_SETTINGS_GROUP_NAME,
                    )
                else:
                    c.value = "{} {} {}".format(
                        _("Total Price"),
                        _("w TVA"),
                        settings.REPANIER_SETTINGS_GROUP_NAME,
                    )
                c.number_format = FORMAT_TEXT
            if col_num == 6:
                c.value = "=" + "+".join(formula_main_total)
                c.number_format = repanier.apps.REPANIER_SETTINGS_CURRENCY_XLSX
                c.font = FONT_BOLD
            current_col += 1
        if hide_column_unit_deposit:
            ws.column_dimensions[get_column_letter(6)].hidden = True
        if not show_column_reference:
            ws.column_dimensions[get_column_letter(3)].hidden = True
    translation.activate(language_code)
    return wb
